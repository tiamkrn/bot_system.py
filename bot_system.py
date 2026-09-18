#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import json
from datetime import datetime, timedelta
from pathlib import Path
import logging

from telegram import Update
from telegram.ext import Application, MessageHandler, CommandHandler, filters, ContextTypes
import pytz
from openpyxl import Workbook

logging.basicConfig(format='%(asctime)s - %(levelname)s - %(message)s', level=logging.INFO)
logger = logging.getLogger(__name__)

with open('config.json', 'r', encoding='utf-8') as f:
    CONFIG = json.load(f)

DATA_DIR = Path('data')
DATA_DIR.mkdir(exist_ok=True)

def load_data(file_path):
    if file_path.exists():
        with open(file_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

def save_data(file_path, data):
    with open(file_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def format_datetime():
    tz = pytz.timezone('Asia/Tehran')
    now = datetime.now(tz)
    return now.strftime('%Y/%m/%d - %H:%M')

def get_week_dates():
    """دوشنبه جاری تا دوشنبه بعدی"""
    today = datetime.now()
    weekday = today.weekday()
    
    if weekday == 0:
        start = today
    else:
        start = today - timedelta(days=weekday)
    
    end = start + timedelta(days=7)
    
    return start.strftime('%Y-%m-%d'), end.strftime('%Y-%m-%d')

def create_excel(start_date, end_date):
    """ایجاد فایل Excel از داده‌ها"""
    
    wb = Workbook()
    
    # ترون
    ws_tron = wb.active
    ws_tron.title = "ترون"
    ws_tron['A1'] = "یوزر"
    ws_tron['B1'] = "مقدار TRX"
    ws_tron['C1'] = "ژتون"
    ws_tron['D1'] = "والت"
    ws_tron['E1'] = "تاریخ"
    
    tron_data = load_data(DATA_DIR / 'tron_data.json')
    row = 2
    for item in tron_data:
        if start_date <= item['timestamp'][:10] <= end_date:
            ws_tron[f'A{row}'] = item['username']
            ws_tron[f'B{row}'] = item['amount']
            ws_tron[f'C{row}'] = item['tokens']
            ws_tron[f'D{row}'] = item['wallet']
            ws_tron[f'E{row}'] = item['timestamp']
            row += 1
    
    # تتر
    ws_tether = wb.create_sheet("تتر")
    ws_tether['A1'] = "یوزر"
    ws_tether['B1'] = "مقدار USDT"
    ws_tether['C1'] = "ژتون"
    ws_tether['D1'] = "والت"
    ws_tether['E1'] = "تاریخ"
    
    tether_data = load_data(DATA_DIR / 'tether_data.json')
    row = 2
    for item in tether_data:
        if start_date <= item['timestamp'][:10] <= end_date:
            ws_tether[f'A{row}'] = item['username']
            ws_tether[f'B{row}'] = item['amount']
            ws_tether[f'C{row}'] = item['tokens']
            ws_tether[f'D{row}'] = item['wallet']
            ws_tether[f'E{row}'] = item['timestamp']
            row += 1
    
    # کش 10+
    ws_cash_plus = wb.create_sheet("کش 10+")
    ws_cash_plus['A1'] = "یوزر"
    ws_cash_plus['B1'] = "مقدار"
    ws_cash_plus['C1'] = "شماره کارت"
    ws_cash_plus['D1'] = "نام صاحب"
    ws_cash_plus['E1'] = "تاریخ"
    
    cash_plus_data = load_data(DATA_DIR / 'cashout_10plus.json')
    row = 2
    for item in cash_plus_data:
        if start_date <= item['timestamp'][:10] <= end_date:
            ws_cash_plus[f'A{row}'] = item['username']
            ws_cash_plus[f'B{row}'] = item['amount']
            ws_cash_plus[f'C{row}'] = item['card']
            ws_cash_plus[f'D{row}'] = item['owner']
            ws_cash_plus[f'E{row}'] = item['timestamp']
            row += 1
    
    # کش 10-
    ws_cash_minus = wb.create_sheet("کش 10-")
    ws_cash_minus['A1'] = "یوزر"
    ws_cash_minus['B1'] = "مقدار"
    ws_cash_minus['C1'] = "شماره کارت"
    ws_cash_minus['D1'] = "نام صاحب"
    ws_cash_minus['E1'] = "تاریخ"
    
    cash_minus_data = load_data(DATA_DIR / 'cashout_10minus.json')
    row = 2
    for item in cash_minus_data:
        if start_date <= item['timestamp'][:10] <= end_date:
            ws_cash_minus[f'A{row}'] = item['username']
            ws_cash_minus[f'B{row}'] = item['amount']
            ws_cash_minus[f'C{row}'] = item['card']
            ws_cash_minus[f'D{row}'] = item['owner']
            ws_cash_minus[f'E{row}'] = item['timestamp']
            row += 1
    
    # کش کریپتو
    ws_crypto = wb.create_sheet("کش کریپتو")
    ws_crypto['A1'] = "یوزر"
    ws_crypto['B1'] = "مقدار"
    ws_crypto['C1'] = "والت"
    ws_crypto['D1'] = "تاریخ"
    
    crypto_data = load_data(DATA_DIR / 'cashout_crypto.json')
    row = 2
    for item in crypto_data:
        if start_date <= item['timestamp'][:10] <= end_date:
            ws_crypto[f'A{row}'] = item['username']
            ws_crypto[f'B{row}'] = item['amount']
            ws_crypto[f'C{row}'] = item['wallet']
            ws_crypto[f'D{row}'] = item['timestamp']
            row += 1
    
    # ذخیره
    filename = f"weekly_report_{start_date}_to_{end_date}.xlsx"
    wb.save(filename)
    return filename

async def excel_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """دستور /excel - هفتگی دوشنبه تا دوشنبه"""
    try:
        start_date, end_date = get_week_dates()
        
        logger.info(f"📊 ایجاد Excel: {start_date} تا {end_date}")
        
        filename = create_excel(start_date, end_date)
        
        with open(filename, 'rb') as f:
            await update.message.reply_document(f, caption=f"📊 گزارش هفتگی\n{start_date} تا {end_date}")
        
        logger.info(f"✅ Excel ارسال شد: {filename}")
    
    except Exception as e:
        logger.error(f"❌ خطا: {e}")
        await update.message.reply_text(f"❌ خطا: {e}")

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        if not update.message or (update.message.text and update.message.text.startswith('/')):
            return
        
        text = update.message.text or update.message.caption
        if not text:
            return
        
        thread_id = update.message.message_thread_id if update.message.is_topic_message else None
        
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        
        logger.info(f"📨 پیام دریافت - Thread: {thread_id}, خطوط: {len(lines)}")
        
        if len(lines) < 3:
            logger.warning(f"❌ خطوط کافی نیست: {len(lines)}")
            return
        
        # ترون - Thread 3
        if thread_id == 3:
            username = lines[0]
            amount = lines[1]
            tokens = lines[2]
            wallet = lines[3] if len(lines) > 3 else "نامشخص"
            
            data = load_data(DATA_DIR / 'tron_data.json')
            data.append({'username': username, 'amount': amount, 'tokens': tokens, 'wallet': wallet, 'timestamp': format_datetime()})
            save_data(DATA_DIR / 'tron_data.json', data)
            
            response = f"✅ **ترون**\n👤 {username}\n💰 {amount}\n🪙 {tokens}\n🏦 {wallet}"
            await update.message.reply_text(response, parse_mode='Markdown')
            await update.message.delete()
            logger.info(f"✅ ترون: {username}")
        
        # تتر - Thread 2
        elif thread_id == 2:
            username = lines[0]
            amount = lines[1]
            tokens = lines[2]
            wallet = lines[3] if len(lines) > 3 else "نامشخص"
            
            data = load_data(DATA_DIR / 'tether_data.json')
            data.append({'username': username, 'amount': amount, 'tokens': tokens, 'wallet': wallet, 'timestamp': format_datetime()})
            save_data(DATA_DIR / 'tether_data.json', data)
            
            response = f"✅ **تتر**\n👤 {username}\n💰 {amount}\n🪙 {tokens}\n🏦 {wallet}"
            await update.message.reply_text(response, parse_mode='Markdown')
            await update.message.delete()
            logger.info(f"✅ تتر: {username}")
        
        # کش 10+ - Thread 6
        elif thread_id == 6:
            username = lines[0]
            amount = lines[1]
            card = lines[2]
            owner = lines[3] if len(lines) > 3 else "نامشخص"
            
            data = load_data(DATA_DIR / 'cashout_10plus.json')
            data.append({'username': username, 'amount': amount, 'card': card, 'owner': owner, 'timestamp': format_datetime()})
            save_data(DATA_DIR / 'cashout_10plus.json', data)
            
            response = f"✅ **کش 10+**\n👤 {username}\n💰 {amount}\n💳 {card}\n📝 {owner}"
            await update.message.reply_text(response, parse_mode='Markdown')
            await update.message.delete()
            logger.info(f"✅ کش 10+: {username}")
        
        # کش 10- - Thread 7
        elif thread_id == 7:
            username = lines[0]
            amount = lines[1]
            card = lines[2]
            owner = lines[3] if len(lines) > 3 else "نامشخص"
            
            data = load_data(DATA_DIR / 'cashout_10minus.json')
            data.append({'username': username, 'amount': amount, 'card': card, 'owner': owner, 'timestamp': format_datetime()})
            save_data(DATA_DIR / 'cashout_10minus.json', data)
            
            response = f"✅ **کش 10-**\n👤 {username}\n💰 {amount}\n💳 {card}\n📝 {owner}"
            await update.message.reply_text(response, parse_mode='Markdown')
            await update.message.delete()
            logger.info(f"✅ کش 10-: {username}")
        
        # کش کریپتو - Thread 5
        elif thread_id == 5:
            username = lines[0]
            amount = lines[1]
            wallet = lines[2]
            
            data = load_data(DATA_DIR / 'cashout_crypto.json')
            data.append({'username': username, 'amount': amount, 'wallet': wallet, 'timestamp': format_datetime()})
            save_data(DATA_DIR / 'cashout_crypto.json', data)
            
            response = f"✅ **کش کریپتو**\n👤 {username}\n💰 {amount}\n📍 {wallet}"
            await update.message.reply_text(response, parse_mode='Markdown')
            await update.message.delete()
            logger.info(f"✅ کش کریپتو: {username}")
        
        else:
            logger.warning(f"⚠️ Thread ID نامشخص: {thread_id}")
    
    except Exception as e:
        logger.error(f"❌ خطا: {e}")

def main():
    logger.info("🚀 شروع سیستم 5 بات...")
    app = Application.builder().token(CONFIG['bots']['tron']['token']).build()
    
    app.add_handler(CommandHandler("excel", excel_command))
    app.add_handler(MessageHandler(filters.TEXT | filters.PHOTO, handle_message))
    
    logger.info("✅ آماده!")
    app.run_polling(allowed_updates=['message'])

if __name__ == '__main__':
    main()
